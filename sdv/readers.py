"""Structural read: turn an attachment (txt / pdf / docx / xlsx) into plain lines of text.

Every reader returns the same thing, DocText, so extraction never cares about the file format.
Table-shaped formats (docx tables, xlsx sheets) are flattened to "Label: first line" plus indented
continuation lines, which is the same shape the plain-text documents already have.

A file that cannot be parsed, or a PDF with no text layer (a scan), is reported as not readable. The
pipeline escalates those to a person; it never guesses from them.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional


@dataclass
class DocText:
    kind: str  # txt | pdf | docx | xlsx
    lines: list = field(default_factory=list)
    readable: bool = True
    error: Optional[str] = None
    text_layer: bool = True  # False for image-only PDFs
    pages: int = 1


def _clean(s: str) -> str:
    return s.replace(" ", " ").replace("\r", "")


def read_text_bytes(data: bytes) -> DocText:
    text = data.decode("utf-8", errors="replace")
    lines = [_clean(l).rstrip() for l in text.splitlines()]
    if not "".join(lines).strip():
        return DocText("txt", lines, readable=False, error="empty file", text_layer=False)
    return DocText("txt", lines)


def _rotated_text(page) -> str:
    """Text of a page whose /Rotate is 90/180/270. pdfplumber reports such text as vertical runs, so put the
    lines back together ourselves: lines are the columns (or rows, at 180), read in the page's visual order."""
    rot = page.rotation
    chars = [c for c in page.chars if c["text"] != "\n"]
    if not chars:
        return ""
    # (line coordinate, position along the line, line order sign, position order sign)
    if rot == 90:
        line_of, pos_of, line_rev, pos_rev = (lambda c: c["x0"]), (lambda c: c["top"]), True, False
    elif rot == 270:
        line_of, pos_of, line_rev, pos_rev = (lambda c: c["x0"]), (lambda c: c["top"]), False, True
    else:  # 180
        line_of, pos_of, line_rev, pos_rev = (lambda c: c["top"]), (lambda c: c["x0"]), True, True
    chars.sort(key=line_of)
    lines, cur, last = [], [], None
    for c in chars:
        k = line_of(c)
        if last is not None and abs(k - last) > 3:
            lines.append(cur)
            cur = []
        cur.append(c)
        last = k
    lines.append(cur)
    if line_rev:
        lines.reverse()
    out = []
    for ln in lines:
        ln.sort(key=pos_of, reverse=pos_rev)
        txt, prev = "", None
        gaps = [abs(pos_of(b) - pos_of(a)) for a, b in zip(ln, ln[1:])]
        typical = sorted(gaps)[len(gaps) // 2] if gaps else 0
        for c in ln:
            if prev is not None and typical and abs(pos_of(c) - pos_of(prev)) > 1.9 * typical and not txt.endswith(" ") and c["text"] != " ":
                txt += " "
            txt += c["text"]
            prev = c
        out.append(txt.rstrip())
    return "\n".join(out)


def read_pdf_bytes(data: bytes) -> DocText:
    try:
        import pdfplumber
    except ImportError as e:  # environment problem, not a property of the document
        from .deps import MissingDependency

        raise MissingDependency("pdfplumber is not installed; run: python -m pip install -r requirements.txt") from e
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            # use_text_flow keeps each text run intact: without it a long label that overlaps the value column
            # gets interleaved with the value ("ConsKiTgPne CeO." for "Consignee" + "KTP CO.").
            chunks = [(_rotated_text(p) if p.rotation in (90, 180, 270) else (p.extract_text(use_text_flow=True, layout=True) or ""))
                      for p in pdf.pages]
    except Exception as e:  # corrupt / truncated / not a PDF: nothing to read
        return DocText("pdf", readable=False, error=f"cannot be parsed ({type(e).__name__}: {str(e)[:80]})", text_layer=False)
    lines = []
    for chunk in chunks:
        lines.extend(_clean(l).rstrip() for l in chunk.splitlines())
    while lines and not lines[0].strip():
        lines.pop(0)
    if not "".join(lines).strip():
        return DocText("pdf", [], readable=False, error="no text layer (scanned / image-only)", text_layer=False, pages=pages)
    return DocText("pdf", lines, pages=pages)


def _cell_lines(cell_text: str) -> list:
    return [l.strip() for l in _clean(cell_text).split("\n") if l.strip()]


def _emit(label: str, value_lines: list) -> list:
    """label + value lines -> 'Label: first' followed by indented continuation lines."""
    label = label.strip()
    if not value_lines:
        return [f"{label}:"]
    out = [f"{label}: {value_lines[0]}"]
    out.extend("    " + v for v in value_lines[1:])
    return out


def read_docx_bytes(data: bytes) -> DocText:
    try:
        import docx
    except ImportError as e:  # environment problem, not a property of the document
        from .deps import MissingDependency

        raise MissingDependency("python-docx is not installed; run: python -m pip install -r requirements.txt") from e
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        return DocText("docx", readable=False, error=f"cannot be parsed ({type(e).__name__})", text_layer=False)
    lines = []
    # Body order matters for typing (title first), so walk paragraphs then tables.
    for p in d.paragraphs:
        if p.text.strip():
            lines.append(_clean(p.text).strip())
    for t in d.tables:
        for row in t.rows:
            cells = [c.text for c in row.cells]
            # de-duplicate merged cells
            uniq = []
            for c in cells:
                if not uniq or c != uniq[-1]:
                    uniq.append(c)
            if len(uniq) >= 2:
                lines.extend(_emit(_cell_lines(uniq[0])[0] if _cell_lines(uniq[0]) else "", [x for c in uniq[1:] for x in _cell_lines(c)]))
            elif len(uniq) == 1 and uniq[0].strip():
                lines.extend(_cell_lines(uniq[0]))
    if not "".join(lines).strip():
        return DocText("docx", lines, readable=False, error="empty document", text_layer=False)
    return DocText("docx", lines)


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_xlsx_bytes(data: bytes) -> DocText:
    try:
        import openpyxl
    except ImportError as e:  # environment problem, not a property of the document
        from .deps import MissingDependency

        raise MissingDependency("openpyxl is not installed; run: python -m pip install -r requirements.txt") from e
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        return DocText("xlsx", readable=False, error=f"cannot be parsed ({type(e).__name__})", text_layer=False)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[sheet] {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [_fmt_cell(c).strip() for c in row]
            if not any(cells):
                continue
            if len(cells) >= 2 and cells[0] and (cells[1] or len(cells) == 2):
                # "First line | address line ; address line" -> value first line, then continuation
                parts = [p.strip() for p in cells[1].split(" | ") if p.strip()] if cells[1] else []
                lines.extend(_emit(cells[0], parts))
            else:
                lines.append(" ".join(c for c in cells if c))
    if len(lines) <= 1:
        return DocText("xlsx", lines, readable=False, error="empty workbook", text_layer=False)
    return DocText("xlsx", lines)


def read_document(name: str, data: bytes) -> DocText:
    """Dispatch on the file extension (format only; document TYPE is decided later, from content)."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    # Sniff bytes too: a mislabelled file should still be read as what it is.
    if data[:5] == b"%PDF-" or ext == "pdf":
        return read_pdf_bytes(data)
    if data[:2] == b"PK":
        if ext == "xlsx":
            return read_xlsx_bytes(data)
        return read_docx_bytes(data)
    if ext == "docx":
        return read_docx_bytes(data)
    if ext == "xlsx":
        return read_xlsx_bytes(data)
    return read_text_bytes(data)
